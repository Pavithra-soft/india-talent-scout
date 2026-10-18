import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth import COOKIE_NAME, create_session_token
from app.database import get_db
from app.jd_rank import parse_job_description, score_profile
from app.jd_store import filter_profiles
from app.main import app
from app.models import Base, JdProfile
from app.profile_fields import find_notice_period, normalize_phone
from app.profile_parse import parse_upload
from app.excel_export import profiles_to_xlsx

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def test_indian_phone_formats():
    assert normalize_phone("+91 98765 43210") == "+919876543210"
    assert normalize_phone("98765-43210") == "+919876543210"
    assert normalize_phone("09876543210") == "+919876543210"
    assert normalize_phone("91 9876543210") == "+919876543210"
    assert normalize_phone("12345") == ""
    assert normalize_phone("") == ""


def test_notice_period_normalization():
    assert find_notice_period("Notice period: immediate")[1] == 0
    assert find_notice_period("Notice period: 30 days")[1] == 30
    assert find_notice_period("Notice period: 2 months")[1] == 60
    raw, days = find_notice_period("Notice period: serving notice")
    assert "serving notice" in raw.lower()
    assert days is None


def test_skill_synonyms_and_score():
    job = parse_job_description(SAMPLES.joinpath("jd.txt").read_text())
    assert "java" in job.must_have
    assert "javascript" in job.nice_to_have
    assert "kubernetes" in job.nice_to_have
    assert job.min_years == 5 and job.max_years == 9
    assert "Bengaluru" in job.locations

    strong = score_profile("Java, Spring Boot, Kafka, JS", 8, "Bengaluru", job)
    weak = score_profile("Python, Django", 4, "Pune", job)
    assert "javascript" in strong.matching
    assert strong.score > weak.score
    assert "kafka" in weak.missing
    assert strong.rationale


def test_parse_sample_resume_leaves_blanks_blank():
    content = (SAMPLES / "resumes" / "meera-iyer.txt").read_bytes()
    result = parse_upload("meera-iyer.txt", content)
    assert result.errors == []
    profile = result.profiles[0]
    assert profile.email == "meera.iyer@example.com"
    assert profile.phone == ""
    assert profile.notice_days == 60
    assert profile.location == "Bengaluru"


def test_filters_combine_with_and():
    rows = [
        JdProfile(name="Ananya Rao", email="a@example.com", phone="+919876543210", location="Bengaluru",
                  skills="java", matching_skills="java", missing_skills="kafka", company="Infosys",
                  notice_days=30, score=80),
        JdProfile(name="Rahul Menon", email="r@example.com", phone="+919876543211", location="Pune",
                  skills="python", matching_skills="", missing_skills="java", company="Tcs",
                  notice_days=0, score=20),
    ]
    filtered = filter_profiles(rows, {"notice": "30", "locations": ["Bengaluru"], "name": "Ananya"})
    assert [row.name for row in filtered] == ["Ananya Rao"]
    assert filter_profiles(rows, {"notice": "30", "locations": ["Pune"], "name": "Ananya"}) == []


def test_excel_headers_and_autofilter():
    content = profiles_to_xlsx(
        [JdProfile(name="Ananya Rao", email="a@example.com", phone="+919876543210", location="Bengaluru",
                   skills="java", matching_skills="java", missing_skills="", company="Infosys",
                   notice_raw="30 days", notice_days=30, score=90, rationale="Covers skills.")]
    )
    sheet = load_workbook(io.BytesIO(content)).active
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert headers[:4] == ["Rank", "Name", "Email", "Phone"]
    assert headers[-1] == "Match Score"
    assert sheet.auto_filter.ref.startswith("A1:")
    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value == 1


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)

    def override():
        db = session_factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override
    with TestClient(app) as test_client:
        test_client.cookies.set(COOKIE_NAME, create_session_token())
        yield test_client
    app.dependency_overrides.clear()


def test_upload_rank_export_flow(client):
    files = []
    for path in sorted((SAMPLES / "resumes").glob("*.txt")):
        files.append(("files", (path.name, path.read_bytes(), "text/plain")))
    uploaded = client.post("/recruit/jd/upload", files=files, data={"replace": "0", "confirm_replace": "0"})
    assert uploaded.status_code == 200
    assert uploaded.json()["added"] == 3

    ranked = client.post(
        "/recruit/jd/rank",
        data={"jd_text": SAMPLES.joinpath("jd.txt").read_text(), "confirm_replace": "0",
              "weight_must": "60", "weight_nice": "20", "weight_experience": "10", "weight_location": "10"},
    )
    assert ranked.status_code == 200
    assert ranked.json()["scored"] == 3

    page = client.get("/recruit/jd", params={"locations": "Bengaluru", "notice": "any"})
    assert page.status_code == 200
    assert "Ananya Rao" in page.text
    assert "Rahul Menon" not in page.text

    exported = client.get("/recruit/jd/export.xlsx", params={"scope": "filtered", "locations": "Bengaluru"})
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content)).active
    names = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "Ananya Rao" in names
    assert "Rahul Menon" not in names

    cleared = client.post("/recruit/jd/clear", follow_redirects=False)
    assert cleared.status_code == 303
    assert client.get("/recruit/jd").text.count("profiles loaded")
