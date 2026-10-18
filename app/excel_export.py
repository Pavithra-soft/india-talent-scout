"""One-sheet Excel export of ranked profiles."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import JdProfile

HEADERS = [
    "Rank",
    "Name",
    "Email",
    "Phone",
    "Location",
    "Skill Set",
    "Matching Skills",
    "Non-Matching Skills",
    "Previous Company",
    "Notice Period",
    "Match Score",
]


def profiles_to_xlsx(profiles: list[JdProfile]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ranked profiles"
    sheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    ranked = sorted(profiles, key=lambda row: (-(row.score or 0), (row.name or "").lower()))
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, profile in enumerate(ranked, start=1):
        sheet.append(
            [
                index,
                profile.name or "",
                profile.email or "",
                profile.phone or "",
                profile.location or "",
                profile.skills or "",
                profile.matching_skills or "",
                profile.missing_skills or "",
                profile.company or "",
                _notice(profile),
                profile.score if profile.score is not None else "",
            ]
        )
        for cell in sheet[index + 1]:
            cell.alignment = wrap

    last_row = max(2, sheet.max_row)
    last_col = get_column_letter(len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_col}{last_row}"
    sheet.freeze_panes = "A2"

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(48, max(len(str(column[0].value or "")) + 2, 14))
        sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _notice(profile: JdProfile) -> str:
    if profile.notice_raw:
        return profile.notice_raw
    if profile.notice_days is None:
        return ""
    if profile.notice_days == 0:
        return "immediate"
    return f"{profile.notice_days} days"
