"""Export job applications to Excel or ZIP (spreadsheet + resume files)."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.application_display import application_view
from app.models import JobApplication, JobPosting

HEADERS = [
    "Applied at",
    "Name",
    "Email",
    "Phone",
    "Years experience",
    "Skills",
    "Notice period",
    "Additional notes",
    "Location",
    "Company",
    "Resume filename",
]


def applications_to_xlsx(posting: JobPosting, applications: list[JobApplication]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append([f"Job: {posting.title}"])
    sheet.append([f"Posting id: {posting.id}"])
    sheet.append([])
    sheet.append(HEADERS)
    header_row = sheet.max_row
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    wrap = Alignment(wrap_text=True, vertical="top")
    for application in sorted(applications, key=lambda row: row.applied_at):
        view = application_view(application)
        sheet.append(
            [
                _fmt_dt(application.applied_at),
                view.name,
                view.email,
                view.phone,
                view.years_experience if view.years_experience != "—" else "",
                view.skills if view.skills != "—" else "",
                application.notice_raw or "",
                application.additional_notes or "",
                view.location if view.location != "—" else "",
                view.company if view.company != "—" else "",
                application.resume_original_filename or "",
            ]
        )
        for cell in sheet[sheet.max_row]:
            cell.alignment = wrap

    last_col = get_column_letter(len(HEADERS))
    sheet.auto_filter.ref = f"A{header_row}:{last_col}{sheet.max_row}"
    sheet.freeze_panes = f"A{header_row + 1}"

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(48, max(len(str(column[0].value or "")) + 2, 14))
        sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def applications_to_zip(posting: JobPosting, applications: list[JobApplication]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("applications.xlsx", applications_to_xlsx(posting, applications))
        for application in applications:
            if not application.resume_stored_path:
                continue
            path = Path(application.resume_stored_path)
            if not path.is_file():
                continue
            safe_name = _safe_zip_name(application)
            archive.write(path, arcname=f"resumes/{safe_name}")
    return buffer.getvalue()


def _fmt_dt(value: datetime) -> str:
    if value.tzinfo:
        value = value.astimezone(timezone.utc)
    return value.strftime("%Y-%m-%d %H:%M UTC")


def _safe_zip_name(application: JobApplication) -> str:
    base = application.resume_original_filename or f"application_{application.id}.bin"
    base = base.replace("/", "_").replace("\\", "_")
    return f"{application.id}_{base}"
